import io
import logging

import fitz
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app-old.api import routes
from app-old.main import app


client = TestClient(app)
EXTRACTED_TEXT = (
    "Name: Asha Raman\n"
    "Date of Birth: 02/03/2006\n"
    "Gender: F\n"
    "Father's / Guardian's Name: Raman\n"
    "Mother's Name: Latha\n"
    "Mobile No student: 9876543210\n"
    "Email ID Student: asha@example.test\n"
    "Address for Communication: 10 Lake Road\n"
    "Caste: BC\n"
)


def make_pdf():
    document = fitz.open()
    page = document.new_page()
    page.insert_text((40, 40), "Admission form")
    content = document.tobytes()
    document.close()
    return content


def make_image(format_name):
    output = io.BytesIO()
    Image.new("RGB", (24, 24), "white").save(output, format=format_name)
    return output.getvalue()


def create_batch():
    response = client.post("/api/batches")
    assert response.status_code == 200
    return response.json()["batch_id"]


def test_backend_root_redirects_to_the_next_frontend():
    response = client.get("/", follow_redirects=False)

    assert response.status_code == 307
    assert response.headers["location"] == "http://127.0.0.1:3000"


def test_upload_supported_formats_and_editable_corrections(monkeypatch):
    monkeypatch.setattr(
        routes,
        "extract_with_layout",
        lambda path, timings=None: (EXTRACTED_TEXT, "test", []),
    )
    batch_id = create_batch()
    files = [
        ("files", ("sample.pdf", make_pdf(), "application/pdf")),
        ("files", ("sample.png", make_image("PNG"), "image/png")),
        ("files", ("sample.jpeg", make_image("JPEG"), "image/jpeg")),
    ]

    try:
        response = client.post(f"/api/batches/{batch_id}/upload", files=files)
        assert response.status_code == 200
        result = response.json()
        assert [item["status"] for item in result["results"]] == [
            "review_required",
            "review_required",
            "review_required",
        ]
        document = result["documents"][0]
        assert document["fields"]["Student Name"] == "Asha Raman"
        assert client.get(
            f"/api/batches/{batch_id}/documents/{document['document_id']}/file"
        ).status_code == 200

        patch = client.patch(
            f"/api/batches/{batch_id}/documents/{document['document_id']}/fields/Student%20Name",
            json={"value": "Asha R."},
        )
        assert patch.status_code == 200
        assert patch.json()["status"] == "saved"
        refreshed = client.get(f"/api/batches/{batch_id}").json()
        assert refreshed["documents"][0]["fields"]["Student Name"] == "Asha R."

        workbook_response = client.get(f"/api/batches/{batch_id}/export.xlsx")
        assert workbook_response.status_code == 200
        workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=True)
        assert "Asha R." in workbook.sheetnames
        rows = list(workbook["Asha R."].values)
        assert ("Student Name", "Asha R.") in rows
        workbook.close()

        csv_response = client.get(f"/api/batches/{batch_id}/export.csv")
        assert csv_response.status_code == 200
        csv_text = csv_response.content.decode("utf-8-sig")
        assert "Student Name" in csv_text
        assert "Asha R." in csv_text
    finally:
        client.delete(f"/api/batches/{batch_id}")


def test_upload_rejects_more_than_twenty_files():
    batch_id = create_batch()
    files = [
        ("files", (f"form-{index}.pdf", b"", "application/pdf"))
        for index in range(21)
    ]

    try:
        response = client.post(f"/api/batches/{batch_id}/upload", files=files)
        assert response.status_code == 400
        assert "limit is 20" in response.json()["detail"]
    finally:
        client.delete(f"/api/batches/{batch_id}")


def test_extract_endpoint_persists_fields_for_the_requested_document(
    monkeypatch, caplog
):
    monkeypatch.setattr(
        routes,
        "extract_with_layout",
        lambda path, timings=None: (EXTRACTED_TEXT, "test", []),
    )
    batch_id = create_batch()
    uploaded = client.post(
        f"/api/batches/{batch_id}/upload?extract=false",
        files=[("files", ("sample.pdf", make_pdf(), "application/pdf"))],
    )
    assert uploaded.status_code == 200
    document_id = uploaded.json()["results"][0]["document_id"]

    try:
        caplog.set_level(logging.INFO, logger=routes.__name__)
        extracted = client.post(
            f"/api/batches/{batch_id}/documents/{document_id}/extract"
        )
        assert extracted.status_code == 200
        assert extracted.json()["document_id"] == document_id
        assert extracted.json()["fields"]["Student Name"] == "Asha Raman"
        assert "field_mapping_ms" in extracted.json()["extraction_metrics"]
        assert "persistence_ms" in extracted.json()["extraction_metrics"]
        assert "request_total_ms" in extracted.json()["extraction_metrics"]
        assert "raw_text_length" in extracted.json()["extraction_metrics"]

        refreshed = client.get(
            f"/api/batches/{batch_id}/documents/{document_id}"
        )
        assert refreshed.status_code == 200
        assert refreshed.json()["fields"]["Student Name"] == "Asha Raman"
        assert "persistence_ms" in refreshed.json()["extraction_metrics"]
        assert "mapped_fields=" in caplog.text
        assert "raw_text_length=" in caplog.text
        assert "Asha Raman" not in caplog.text
    finally:
        client.delete(f"/api/batches/{batch_id}")


def test_unsupported_upload_is_reported_without_storing_document():
    batch_id = create_batch()

    try:
        response = client.post(
            f"/api/batches/{batch_id}/upload",
            files=[("files", ("form.txt", b"not a form", "text/plain"))],
        )
        assert response.status_code == 200
        assert response.json()["results"][0]["status"] == "failed"
        assert "Only PDF" in response.json()["results"][0]["error"]
        assert response.json()["documents"] == []
    finally:
        client.delete(f"/api/batches/{batch_id}")


def test_batch_deletion_removes_only_its_own_files(monkeypatch):
    monkeypatch.setattr(
        routes,
        "extract_with_layout",
        lambda path, timings=None: (EXTRACTED_TEXT, "test", []),
    )
    batch_id = create_batch()
    unrelated_id = create_batch()
    response = client.post(
        f"/api/batches/{batch_id}/upload",
        files=[("files", ("sample.pdf", make_pdf(), "application/pdf"))],
    )
    assert response.status_code == 200
    document = response.json()["documents"][0]
    file_response = client.get(
        f"/api/batches/{batch_id}/documents/{document['document_id']}/file"
    )
    assert file_response.status_code == 200

    try:
        deleted = client.delete(f"/api/batches/{batch_id}")
        assert deleted.status_code == 200
        assert client.get(f"/api/batches/{batch_id}").status_code == 404
        assert client.get(f"/api/batches/{unrelated_id}").status_code == 200
        assert client.get(
            f"/api/batches/{batch_id}/documents/{document['document_id']}/file"
        ).status_code == 404
    finally:
        client.delete(f"/api/batches/{unrelated_id}")
