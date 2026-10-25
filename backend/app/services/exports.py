"""Excel and CSV exports for the Admission Form Digitization System.

The XLSX layout intentionally follows the team-lead reference workbook:
header/source lines, Field/Extracted information table, side-by-side XI/XII
marks table, signatures, certificate/office entries, and review notes.
"""

import csv
import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Keep the reference workbook's field labels and order. Values are mapped from
# the application's existing canonical extraction fields.
FIELD_ROWS = [
    ("Application No.", "Application No"),
    ("Roll No.", "Roll No"),
    ("Roll number correction", "Roll number correction"),
    ("Received By", "Received By"),
    ("Received Date", "Received Date"),
    ("Received Date – original text", "Received Date – original text"),
    ("Proposed Study", "Course Applied"),
    ("Name", "Student Name"),
    ("Gender", "Gender"),
    ("Father’s / Guardian’s Name", "Father's Name"),
    ("Mother’s Name", "Mother's Name"),
    ("Religion", "Religion"),
    ("Date of Birth", "Date of Birth"),
    ("Nationality", "Nationality"),
    ("Community", "Community"),
    ("Caste", "Caste"),
    ("Sub Caste", "Sub Caste"),
    ("Aadhaar No.", "Aadhaar No"),
    ("PAN No.", "Pan No"),
    ("Bank Name", "Bank Name"),
    ("Bank A/C No.", "Bank A/C No"),
    ("IFSC Code", "IFSC Code"),
    ("Blood Group", "Blood Group"),
    ("Email ID – Student", "Email"),
    ("Mobile No. – Student", "Mobile Number"),
    ("Student mobile annotation", "Student mobile annotation"),
    ("Father Mobile No.", "Father Mobile No"),
    ("Tel No.", "Telephone No"),
    ("Contact across Father Mobile / Tel lines", "Contact across Father Mobile / Tel lines"),
    ("EMIS No.", "EMIS No"),
    ("Identification Mark 1", "Identification Mark 1"),
    ("Identification Mark 2", "Identification Mark 2"),
    ("Parents Occupation", "Parents Occupation"),
    ("Income (as written, numeric)", "Income (as written, numeric)"),
    ("Income – original notation", "Income – original notation"),
    ("Physically Challenged", "Physically Challenged"),
    ("First Graduate", "First Graduate"),
    ("Admission No.", "Admission No"),
    ("Admission Fee", "Admission Fee"),
    ("Address for Communication", "Address"),
    ("Last Studied Institution with Address", "Last Studied Institution with Address"),
    ("Medium", "Medium"),
]

CERTIFICATE_FIELDS = [
    ("Transfer Certificate (Original)", "Transfer Certificate"),
    ("Mark Sheet — X, XI, XII (Original)", "Mark Sheets"),
    ("Community Certificate (Proof)", "Community Certificate"),
    ("Income Certificate (Proof)", "Income Certificate"),
    ("Passport-size Photo (4)", "Passport-size Photo"),
    ("Aadhaar (Proof)", "Aadhaar Proof"),
    ("PAN (Proof)", "PAN Proof"),
    ("Bank Passbook (Proof)", "Bank Passbook Proof"),
]

SUBJECTS = ["Tamil", "English", "Physics", "Chemistry", "Botany", "Zoology", "Biology", "Mathematics", "Maths"]

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F3F6FA"
GRID = "D7DEE8"
DARK = "243247"
WHITE = "FFFFFF"


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value
    return str(value).strip()


def _value(fields, key, default="Not detected — enter value"):
    """Respect extraction's distinction: '-' means absent; blank means unknown."""
    if key not in fields:
        return default
    value = fields.get(key)
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    return value


def _first(fields, *keys, default="Not detected — enter value"):
    for key in keys:
        if key in fields:
            value = fields.get(key)
            if value is not None and not (isinstance(value, str) and not value.strip()):
                return value
    return default


def _clean_sheet_title(fields, index):
    raw = _first(fields, "Roll No", "Student Name", default=f"Student {index}")
    title = re.sub(r'[\[\]:*?/\\]', "", str(raw)).strip() or f"Student {index}"
    title = title[:31]
    return title


def _unique_title(base, used):
    title = base
    n = 2
    while title in used:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def _parse_subjects(value):
    """Parse 'Tamil: 35, English: 40' into a subject -> mark mapping."""
    if not isinstance(value, str) or not value.strip():
        return {}
    result = {}
    for match in re.finditer(r"([A-Za-z][A-Za-z &./-]*?)\s*:\s*([^,;]+)", value):
        label = re.sub(r"\s+", " ", match.group(1)).strip()
        mark = match.group(2).strip()
        if label.lower() in {"total", "percentage", "maximum", "reg. no", "register no"}:
            continue
        result[label.lower()] = mark
    return result


def _subject_value(records, subject):
    for name, value in records.items():
        if name == subject.lower() or (subject.lower() == "mathematics" and name == "maths") or (subject.lower() == "maths" and name == "mathematics"):
            return value
    return "Not detected — enter value"


def _style_section(ws, row, title):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    cell = ws.cell(row, 1, title)
    cell.font = Font(bold=True, color=WHITE, size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22


def _build_student_sheet(wb, document, index, used_titles):
    fields = document.get("fields") or {}
    roll = _first(fields, "Roll No", default="")
    name = _first(fields, "Student Name", "Name", default=f"Student {index}")
    sheet_base = str(roll if roll and str(roll).strip() not in {"-", "—"} else name)
    ws = wb.create_sheet(_unique_title(_clean_sheet_title({"Student Name": sheet_base}, index), used_titles))
    ws.sheet_view.showGridLines = False

    # Title and source header: same hierarchy as the reference workbook.
    ws.merge_cells("A2:C2")
    ws["A2"] = f"{roll if roll and str(roll).strip() not in {'-', '—'} else f'Student {index}'} — {name}"
    ws["A2"].font = Font(bold=True, size=15, color=WHITE)
    ws["A2"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A2"].alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 30
    ws.merge_cells("A3:C3")
    ws["A3"] = "Application for Admission to U.G. Degree Programs"
    ws["A3"].font = Font(bold=True, size=12, color=DARK)
    ws["A3"].alignment = Alignment(vertical="center")
    ws.merge_cells("A4:C4")
    source = document.get("original_filename") or "source document not recorded"
    ws["A4"] = f"Source: {source}"
    ws["A4"].font = Font(italic=True, size=9, color="5B6575")
    ws["A4"].alignment = Alignment(vertical="center")

    # Reference workbook's field/value section.
    ws.cell(6, 1, "Field")
    ws.cell(6, 2, "Extracted information")
    for c in ws[6][:2]:
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center")
    row = 7
    for label, key in FIELD_ROWS:
        value = _value(fields, key)
        # The reference workbook represents selected checkbox options as the
        # option in the field label and "Yes" in the value column.
        if key == "Course Applied":
            selected = _first(fields, "Course Applied", "Proposed Study", default="Not detected — enter value")
            if str(selected).strip() not in {"-", "—", "Not detected — enter value", ""}:
                label, value = f"Proposed Study – {selected}", "Yes"
            else:
                label, value = "Proposed Study", selected
        elif key in {"Gender", "Religion", "Community", "First Graduate"}:
            selected = _value(fields, key)
            if str(selected).strip() not in {"-", "—", "Not detected — enter value", ""}:
                label = f"{label} – {selected}"
                value = "Yes"
        elif key == "Received Date – original text":
            value = _first(fields, "Received Date – original text", "Received Date", default="Not detected — enter value")
        elif key in {"Roll number correction", "Student mobile annotation", "Contact across Father Mobile / Tel lines"} and key not in fields:
            # These are auxiliary notes rather than canonical extraction fields;
            # don't falsely flag them as failed extraction when not supplied.
            value = "—"
        ws.cell(row, 1, label)
        ws.cell(row, 2, value)
        if row % 2:
            for c in ws[row][:2]:
                c.fill = PatternFill("solid", fgColor=PALE)
        row += 1

    # Academic table with XI and XII kept side-by-side as in the reference.
    row += 2
    for col, value in enumerate(["Subject", "XI Standard", "XII Standard"], 1):
        cell = ws.cell(row, col, value)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    row += 1
    xi = _parse_subjects(_first(fields, "XI Academic Records", default=""))
    xii = _parse_subjects(_first(fields, "XII Academic Records", default=""))
    ordered_subjects = ["Tamil", "English", "Physics", "Chemistry", "Botany", "Zoology"]
    extra_subjects = []
    for records in (xi, xii):
        for subject in records:
            pretty = next((s for s in SUBJECTS if s.lower() == subject or (s.lower() == "mathematics" and subject == "maths")), subject.title())
            if pretty.lower() not in {s.lower() for s in ordered_subjects + extra_subjects}:
                extra_subjects.append(pretty)
    for subject in ordered_subjects + extra_subjects:
        ws.cell(row, 1, subject)
        ws.cell(row, 2, _subject_value(xi, subject))
        ws.cell(row, 3, _subject_value(xii, subject))
        row += 1
    for label, xi_key, xii_key in [
        ("Total marks – written", "XI Standard Marks Obtained", "XII Standard Marks Obtained"),
        ("Percentage – written", "XI Standard Percentage", "XII Standard Percentage"),
    ]:
        ws.cell(row, 1, label)
        ws.cell(row, 2, _first(fields, xi_key, default="Not detected — enter value"))
        ws.cell(row, 3, _first(fields, xii_key, default="Not detected — enter value"))
        row += 1

    row += 1
    _style_section(ws, row, "Signatures / Other Filled Information")
    row += 1
    parent_signature = _first(fields, "Parent / Guardian Signature", default="Not detected — enter value")
    applicant_signature = _first(fields, "Applicant Signature", default="Not detected — enter value")
    signature_text = f"Parent/guardian signature: {parent_signature}; Applicant signature: {applicant_signature}"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.cell(row, 1, signature_text)
    ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    row += 2

    _style_section(ws, row, "Certificates / Office Entries")
    row += 1
    cert_parts = []
    for label, key in CERTIFICATE_FIELDS:
        cert_parts.append(f"{label}: {_value(fields, key)}")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.cell(row, 1, "; ".join(cert_parts))
    ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    row += 2

    _style_section(ws, row, "Review notes")
    row += 1
    notes = _first(fields, "Review notes", "Review Notes", default="—")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.cell(row, 1, notes)
    ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")

    # Polished, readable presentation while retaining the reference structure.
    ws.column_dimensions["A"].width = 43
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 3
    thin = Side(style="thin", color=GRID)
    for row_cells in ws.iter_rows(min_row=6, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row_cells:
            if cell.value is not None:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="center" if cell.column in (2, 3) and cell.row > 6 else "left",
                )
                if cell.row > 6 and cell.row % 2 == 1 and cell.fill.fill_type is None:
                    cell.fill = PatternFill("solid", fgColor=PALE)
    for r in range(7, ws.max_row + 1):
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 19)
    ws.freeze_panes = "A7"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "portrait"
    ws.print_title_rows = "1:6"
    ws.sheet_properties.outlinePr.summaryBelow = True
    return ws


def excel_bytes(batch):
    wb = Workbook()
    wb.remove(wb.active)
    docs = batch.get("documents", [])
    if not docs:
        ws = wb.create_sheet("No students")
        ws["A1"] = "No student documents"
        ws["A1"].font = Font(bold=True, color=WHITE)
        ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
        ws.column_dimensions["A"].width = 30
    used_titles = set()
    for index, document in enumerate(docs, 1):
        _build_student_sheet(wb, document, index, used_titles)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def csv_bytes(batch):
    docs = batch.get("documents", [])
    keys = list(dict.fromkeys(["Document ID", "File Name"] + [key for doc in docs for key in doc.get("fields", {})]))
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=keys, extrasaction="ignore")
    writer.writeheader()
    for document in docs:
        row = {"Document ID": document.get("document_id", ""), "File Name": document.get("original_filename", "")}
        row.update(document.get("fields", {}))
        writer.writerow(row)
    return output.getvalue().encode("utf-8-sig")
