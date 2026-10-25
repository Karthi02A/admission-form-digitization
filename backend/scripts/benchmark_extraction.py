"""Compare extracted fields with the supplied reference workbook."""

import re
import argparse
from datetime import date, datetime
from pathlib import Path
from time import perf_counter

from openpyxl import load_workbook

from app-old.config.settings import ALIASES, CHECKBOX_FIELDS
from app-old.services.extraction import _valid_value, map_fields, normalize_label
from app-old.services.ocr import extract_with_layout, pdf_layout, pdf_text


DOCUMENTS = Path(__file__).resolve().parents[1] / "data" / "documents"
WORKBOOK = DOCUMENTS / "Student_Applications_BFS26001_BFS26010.xlsx"
IGNORED_LABEL_PARTS = (
    "correction",
    "original text",
    "original notation",
    "annotation",
    "across father",
    "review",
    "source",
    "signature",
)
OPTION_REFERENCE_FIELDS = {
    "Course Applied",
    "Gender",
    "First Graduate",
    "Physically Challenged",
}
POSITIVE_SELECTION_MARKERS = {"yes", "true", "checked", "selected", "x", "1"}
NEGATIVE_SELECTION_MARKERS = {"no", "false", "0"}


def _reference_field(label):
    normalized = normalize_label(label)
    if any(part in normalized for part in IGNORED_LABEL_PARTS):
        return None

    matches = [
        (len(alias), field)
        for field, aliases in ALIASES.items()
        for alias in aliases
        if normalized == normalize_label(alias)
        or normalized.startswith(f"{normalize_label(alias)} ")
    ]
    return max(matches)[1] if matches else None


def _reference_value(label, value, field):
    text = str(value).strip()
    if field not in OPTION_REFERENCE_FIELDS:
        return text

    normalized_value = normalize_label(text)
    normalized_label = normalize_label(label)
    aliases = sorted(
        (normalize_label(alias) for alias in ALIASES[field]),
        key=len,
        reverse=True,
    )
    suffix = next(
        (
            normalized_label[len(alias):].strip()
            for alias in aliases
            if normalized_label.startswith(f"{alias} ")
        ),
        "",
    )
    if not suffix:
        return text
    if normalized_value in NEGATIVE_SELECTION_MARKERS:
        return None
    if normalized_value in POSITIVE_SELECTION_MARKERS:
        return suffix
    return text


def _reference_values(rows):
    expected = {}
    ambiguous = set()
    for row in rows:
        label = row[0] if row else None
        value = row[1] if len(row) > 1 else None
        if (
            not isinstance(label, str)
            or value is None
            or not _has_expected_value(value)
        ):
            continue
        field = _reference_field(label)
        if not field:
            continue
        reference_value = _reference_value(label, value, field)
        if reference_value is None:
            continue
        if field in expected and (
            _comparable(expected[field], field)
            != _comparable(reference_value, field)
        ):
            ambiguous.add(field)
        else:
            expected[field] = reference_value
    for field in ambiguous:
        expected.pop(field, None)
    return expected, ambiguous


def _comparable(value, field=None):
    if field == "Gender":
        normalized = normalize_label(str(value or ""))
        if normalized in {"m", "male"}:
            return "male"
        if normalized in {"f", "female"}:
            return "female"
    if field in {"Course Applied", "Community"}:
        return normalize_label(str(value or ""))
    if field and "date" in field.casefold():
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        text = str(value or "").strip()
        for date_format in (
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d.%m.%Y",
            "%d-%m-%Y",
        ):
            try:
                return datetime.strptime(text, date_format).date().isoformat()
            except ValueError:
                continue
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _has_expected_value(value):
    comparable = _comparable(value)
    return bool(re.search(r"[a-z0-9]", comparable)) and comparable not in {
        "na",
        "n a",
        "n/a",
        "none",
        "not applicable",
        "not available",
        "not provided",
    }


def _contains_value(text, value, field=None):
    expected = _comparable(value, field)
    source = _comparable(text, field)
    return bool(expected) and expected in source


def _mismatch_cause(field, expected, mapped, embedded_text, ocr_text):
    expected_in_embedded = _contains_value(embedded_text, expected, field)
    expected_in_ocr = _contains_value(ocr_text, expected, field)
    if field in CHECKBOX_FIELDS and expected_in_ocr:
        return "checkbox/layout selection"
    if expected_in_embedded and not expected_in_ocr:
        return "OCR recognition/coverage"
    if expected_in_ocr:
        return "field mapping/layout"
    if not _valid_value(field, expected):
        return "validation"
    if mapped:
        return "OCR recognition/selection"
    return "unavailable/unreadable source or OCR recognition"


def _memory_usage_mib():
    try:
        import psutil
    except ImportError:
        return None, None
    memory = psutil.Process().memory_info()
    peak = getattr(memory, "peak_wset", None)
    return round(memory.rss / (1024 * 1024), 1), (
        round(peak / (1024 * 1024), 1) if peak is not None else None
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--student",
        choices=sorted(path.stem for path in DOCUMENTS.glob("BFS260*.pdf")),
        help="Benchmark one reference form instead of all ten",
    )
    parser.add_argument(
        "--text-only",
        action="store_true",
        help="Use the embedded PDF text layer instead of running OCR",
    )
    parser.add_argument(
        "--diagnostics",
        action="store_true",
        help="Report redacted field mismatch categories and stage timings",
    )
    args = parser.parse_args()
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    totals = {field: {"compared": 0, "matched": 0, "nonempty": 0} for field in ALIASES}
    student_ids = [args.student] if args.student else workbook.sheetnames
    started = perf_counter()

    for student_id in student_ids:
        document_started = perf_counter()
        pdf_path = DOCUMENTS / f"{student_id}.pdf"
        timings = {}
        if args.text_only:
            text_started = perf_counter()
            text = pdf_text(pdf_path)
            timings["pdf_text_ms"] = round((perf_counter() - text_started) * 1000)
            layout_started = perf_counter()
            layout = pdf_layout(pdf_path)
            timings["layout_processing_ms"] = round(
                (perf_counter() - layout_started) * 1000
            )
            method = "pdf_text"
        else:
            text, method, layout = extract_with_layout(pdf_path, timings)
        embedded_text = pdf_text(pdf_path) if args.diagnostics and not args.text_only else text
        embedded_layout = (
            pdf_layout(pdf_path) if args.diagnostics and not args.text_only else layout
        )
        embedded_fields = (
            map_fields(embedded_text, embedded_layout)[0]
            if args.diagnostics and not args.text_only
            else {}
        )
        mapping_started = perf_counter()
        extracted, quality = map_fields(text, layout)
        timings["field_mapping_ms"] = round(
            (perf_counter() - mapping_started) * 1000
        )
        expected, ambiguous_expected = _reference_values(
            workbook[student_id].iter_rows(values_only=True)
        )

        compared = 0
        matched = 0
        missing = 0
        incorrect = 0
        for field, expected_value in expected.items():
            totals[field]["compared"] += 1
            compared += 1
            if extracted.get(field):
                totals[field]["nonempty"] += 1
            if (
                _comparable(extracted.get(field), field)
                == _comparable(expected_value, field)
            ):
                totals[field]["matched"] += 1
                matched += 1
            else:
                if extracted.get(field):
                    incorrect += 1
                else:
                    missing += 1
                if args.diagnostics:
                    direct_match = (
                        _comparable(embedded_fields.get(field), field)
                        == _comparable(expected_value, field)
                    )
                    cause = (
                        "OCR recognition/selection"
                        if direct_match
                        else _mismatch_cause(
                            field,
                            expected_value,
                            extracted.get(field),
                            embedded_text,
                            text,
                        )
                    )
                    print(
                        f"MISMATCH {student_id} | {field} | "
                        f"status={'different' if extracted.get(field) else 'missing'} | "
                        f"cause={cause}",
                        flush=True,
                    )

        rss_mib, peak_rss_mib = _memory_usage_mib()
        extraction_stages = (
            "model_initialization_ms",
            "pdf_inspection_ms",
            "pdf_text_ms",
            "ocr_detection_ms",
            "pdf_render_ms",
            "ocr_inference_ms",
            "layout_processing_ms",
            "checkbox_detection_ms",
            "total_extraction_ms",
            "field_mapping_ms",
        )
        stage_summary = " ".join(
            f"{stage}={timings.get(stage, 0)}"
            for stage in extraction_stages
        )
        print(
            f"{student_id}: method={method}, "
            f"mapped={quality['filled_fields']}/{quality['total_fields']}, "
            f"reference_fields={compared}, exact_matches={matched}, "
            f"nonempty={compared - missing}/{compared}, "
            f"incorrect={incorrect}, missing={missing}, "
            f"ambiguous_reference_fields={len(ambiguous_expected)}, "
            f"rss_mib={rss_mib if rss_mib is not None else 'unavailable'}, "
            f"peak_rss_mib={peak_rss_mib if peak_rss_mib is not None else 'unavailable'}, "
            f"stage_ms=[{stage_summary}], "
            f"seconds={perf_counter() - document_started:.1f}",
            flush=True,
        )

    workbook.close()
    print("\nField accuracy (values are not printed):")
    for field, result in totals.items():
        if not result["compared"]:
            continue
        print(
            f"{field}: exact={result['matched']}/{result['compared']}, "
            f"nonempty={result['nonempty']}/{result['compared']}"
        )

    compared = sum(result["compared"] for result in totals.values())
    matched = sum(result["matched"] for result in totals.values())
    nonempty = sum(result["nonempty"] for result in totals.values())
    accuracy = 100 * matched / compared if compared else 0
    print(
        f"\nOverall exact match: {matched}/{compared} ({accuracy:.1f}%), "
        f"nonempty outputs: {nonempty}/{compared}, "
        f"elapsed_seconds={perf_counter() - started:.1f}"
    )


if __name__ == "__main__":
    main()
